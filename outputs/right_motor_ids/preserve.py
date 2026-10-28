from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
import xml.etree.ElementTree as ET
import openpyxl

out = Path(__file__).parent / '左右臂电机ID备忘.xlsx'
source = Path(r'C:\Users\rhche\Desktop\左臂电机ID备忘.xlsx')
rows = list(openpyxl.load_workbook(out)['全身参考'].values)[-3:]
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
tag = lambda name: '{'+ns['s']+'}'+name
with ZipFile(source) as z:
    files = {n:z.read(n) for n in z.namelist()}
root = ET.fromstring(files['xl/worksheets/sheet2.xml'])
data = root.find('s:sheetData',ns)
for idx, values in enumerate(rows, 28):
    template = deepcopy(data.find(f"s:row[@r='{idx-3}']",ns))
    template.set('r', str(idx))
    for col, (cell,value) in enumerate(zip(template, values),1):
        cell.set('r', f'{chr(64+col)}{idx}')
        for child in list(cell): cell.remove(child)
        if isinstance(value,str):
            cell.set('t','inlineStr')
            ET.SubElement(ET.SubElement(cell,tag('is')),tag('t')).text=value
        else:
            cell.set('t','n')
            ET.SubElement(cell,tag('v')).text=str(value)
    data.append(template)
root.find('s:dimension',ns).set('ref','A1:H30')
af=root.find('s:autoFilter',ns)
if af is not None: af.set('ref','A1:H30')
files['xl/worksheets/sheet2.xml'] = ET.tostring(root,encoding='utf-8',xml_declaration=True)
with ZipFile(out,'w',ZIP_DEFLATED) as z:
    for name,content in files.items(): z.writestr(name,content)
a=openpyxl.load_workbook(source); b=openpyxl.load_workbook(out)
for s in a:
    for row in s:
        for c in row:
            d=b[s.title][c.coordinate]
            assert c.value == d.value, (s.title,c.coordinate)
            assert c._style == d._style
assert list(b['全身参考'].values)[-3:] == rows
assert len(set(r[3] for r in list(b['全身参考'].values)[1:]))==29
assert len(set(r[5] for r in list(b['全身参考'].values)[1:]))==29
print('Verified: original cells and styles unchanged; three new ID pairs saved.')
